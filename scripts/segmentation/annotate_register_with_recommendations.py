#!/usr/bin/env python3
"""
annotate_register_with_recommendations.py — put the recommendations next to
the rows they are about.

`review_flags` (column N) already says *what is wrong* with a row:
`fused_entry:description_reference_run` and so on. It does not say *what to do
about it*. That is what `suggest_row_corrections.py` works out, and this script
writes it into columns O onward so the two sit side by side in one sheet.

Columns A–N are copied through untouched. Nothing is corrected here.

    O  15_fix_kind            which defect, in this script's vocabulary
    P  16_fix_tier            1 mechanical · 2 safe · 3 needs a decision · 4 never automatic
    Q  17_fix_action          what to do
    R  18_recover_refs        citations stranded in prose, not in column K/L
    S  19_proposed_text       what the fused column should be trimmed to
    T  20_split_new_entry     the entry hiding in the tail, if it has no row yet
    U  21_split_existing      the entry hiding in the tail that ALREADY has a row
    V  22_merge_candidate     the other row this one may duplicate
    W  23_evidence            the text the finding was made on

Usage:
    python scripts/segmentation/annotate_register_with_recommendations.py \\
        --register <merged.tsv> --out <annotated>
"""

import argparse
import csv
import importlib.util
import sys
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8")
csv.field_size_limit(10 ** 7)

ROOT = Path(__file__).resolve().parents[2]
_spec = importlib.util.spec_from_file_location(
    "sugg", Path(__file__).resolve().parent / "suggest_row_corrections.py")
sugg = importlib.util.module_from_spec(_spec)
sys.modules["sugg"] = sugg
_spec.loader.exec_module(sugg)

NEW_COLS = ["15_fix_kind", "16_fix_tier", "17_fix_action", "18_recover_refs",
            "19_proposed_text", "20_split_new_entry", "21_split_existing",
            "22_merge_candidate", "23_evidence"]


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[1])
    ap.add_argument("--register", type=Path, required=True)
    ap.add_argument("--out", type=Path, required=True)
    ap.add_argument("--xlsx", action="store_true", help="also write an .xlsx")
    args = ap.parse_args()

    with args.register.open(encoding="utf-8", newline="") as f:
        rdr = csv.DictReader(f, delimiter="\t")
        fields = list(rdr.fieldnames)
        rows = [r for r in rdr if r.get("01_entry_id")]
    print(f"  læst {len(rows):,} rækker, {len(fields)} kolonner (A-"
          f"{chr(64 + len(fields))})")

    # the detector keys on 00_person_id; this register keys on 01_entry_id
    id_col = "00_person_id" if "00_person_id" in fields else "01_entry_id"
    shim = [dict(r, **{"00_person_id": r[id_col]}) for r in rows]
    surnames = {(r["03_surname"] or "").strip().lower() for r in rows}

    findings = sugg.find_splits(shim, surnames) + sugg.find_merges(shim)
    print(f"  {len(findings):,} fund på {len({f['person_id'] for f in findings}):,} rækker")

    ann = {}
    for f in findings:
        for pid in f["person_id"].split():          # merges name two rows
            a = ann.setdefault(pid, {c: [] for c in NEW_COLS})
            a["15_fix_kind"].append(f["kind"])
            a["16_fix_tier"].append(f["tier"])
            a["17_fix_action"].append(f["action"])
            a["23_evidence"].append(f["evidence"])
            act, prop = f["action"], f["proposal"]
            if act == "recover citations":
                a["18_recover_refs"].append(prop)
            elif act == "trim the column":
                a["19_proposed_text"].append(f'{f["column"]} = {prop}')
            elif act == "split out a new row":
                a["20_split_new_entry"].append(prop.replace("new entry: ", ""))
            elif "already exists" in act:
                a["21_split_existing"].append(prop.replace("see existing ", ""))
            elif "merge" in act:
                other = [p for p in f["person_id"].split() if p != pid]
                a["22_merge_candidate"].append(
                    f'{" ".join(other)} — {f["evidence"]} — {prop}')

    out_fields = fields + NEW_COLS
    out_rows = []
    for r in rows:
        a = ann.get(r[id_col], {})
        # de-duplicate while keeping order; one row can carry several findings
        extra = {}
        for c in NEW_COLS:
            vals, seen = [], set()
            for v in a.get(c, []):
                if v and v not in seen:
                    seen.add(v)
                    vals.append(v)
            extra[c] = " | ".join(vals)
        out_rows.append(dict(r, **extra))

    args.out.parent.mkdir(parents=True, exist_ok=True)
    with args.out.open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=out_fields, delimiter="\t",
                           extrasaction="ignore")
        w.writeheader()
        w.writerows(out_rows)
    filled = sum(1 for r in out_rows if r["17_fix_action"])
    print(f"  skrev {args.out}  ({len(out_rows):,} rækker, {filled:,} med anbefaling)")
    for c in NEW_COLS:
        n = sum(1 for r in out_rows if r[c])
        col = chr(65 + out_fields.index(c)) if out_fields.index(c) < 26 else "?"
        print(f"     {col}  {c:20} {n:>5}")

    if args.xlsx:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        wb = Workbook()
        ws = wb.active
        ws.title = "personregister XI"
        ws.append(out_fields)
        for c in range(1, len(out_fields) + 1):
            ws.cell(1, c).font = Font(bold=True)
            if out_fields[c - 1] in NEW_COLS:
                ws.cell(1, c).fill = PatternFill("solid", fgColor="FFE8CC")
        for r in out_rows:
            ws.append([r.get(c, "") for c in out_fields])
        ws.freeze_panes = "A2"
        xp = args.out.with_suffix(".xlsx")
        wb.save(xp)
        print(f"  skrev {xp}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
