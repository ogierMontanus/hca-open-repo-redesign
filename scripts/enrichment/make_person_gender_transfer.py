#!/usr/bin/env python3
"""
make_person_gender_transfer.py
------------------------------
Bygger et overførselsark med de seneste kønsbestemmelser, nationalitet og rolle til
data/raw/6-PERSON-V0.97(5)(1).xlsx (ark "Person", kolonnerne H "Nationality", I "Gender" og
J "Role/Occupation", som er tomme i V0.97).

Rækkerne står i samme rækkefølge som V0.97, så kolonnerne D:F kan kopieres
direkte ind i kolonne H:J fra række 5. PersonID og OLDPersonID følger med som
kontrol.

Kønskilden er det gennemsete kønsgennemsyn
data/review/gender_inference/hca-personregister-redesign_gender-review.xlsx
(alle manuelle afgørelser fra data/curated/gender_manual_review.csv er
indarbejdet; ingen poster er »Endnu ubestemt«).

V0.97 og kønsgennemsynet bruger forskellige id'er og forskellig segmentering
(V0.97: 9.519 poster fra den nuværende hjemmeside; gennemsynet: 10.079 poster
fra den manuelt rettede segmentering). Koblingen går via hjemmesidens
Reg-id'er:

    V0.97  --(label + beskrivelse, ellers rækkefølge)-->  data/normalized/entities.csv (Reg-id)
           --(Reg-ID i gennemsynet, ellers person_id_crosswalk.csv)-->  HCAP-id

Kolonnen "Kobling" viser, hvordan hver række blev koblet; "Tjek" er udfyldt,
hvor koblingen bør ses efter før overførsel.

Kør:
    python scripts/enrichment/make_person_gender_transfer.py
"""

import csv
import difflib
import os
import re
import sys
from collections import Counter, defaultdict

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding="utf-8")
ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
TARGET = os.path.join(ROOT, "data", "raw", "6-PERSON-V0.97(5)(1).xlsx")
REVIEW = os.path.join(ROOT, "data", "review", "gender_inference",
                      "hca-personregister-redesign_gender-review.xlsx")
ENTITIES = os.path.join(ROOT, "data", "normalized", "entities.csv")
CROSSWALK = os.path.join(ROOT, "data", "curated", "person_id_crosswalk.csv")
OUT = os.path.join(ROOT, "data", "review", "gender_inference",
                   "6-PERSON-V0.97_gender-transfer.xlsx")

ORDER_MIN_SCORE = 0.90  # under denne værdi markeres en rækkefølgekobling til tjek
NAME_MIN_SCORE = 0.85   # mindste ligheden mellem V0.97-label og gennemsynets navn, for at en kobling godtages


def name_core(s):
    """Navn uden parenteser (årstal m.m.) og tegnsætning, til sammenligning på tværs af registrene."""
    s = (s or "").replace("–", "-").replace("—", "-")
    s = re.sub(r"\([^)]*\)", "", s)
    return re.sub(r"[\s.,\-]+", " ", s).lower().strip()


def name_sim(a, b):
    return difflib.SequenceMatcher(None, name_core(a), name_core(b)).ratio()


def norm(s):
    s = (s or "").replace("–", "-").replace("—", "-")
    return re.sub(r"[\s.,]+", " ", s).lower().strip()


def fuzzy_match(t, by_first):
    """Bedste kandidat blandt gennemsynets poster med samme første navneord; kræver både navn og beskrivelse tæt på."""
    best, second = (0, None), 0
    for r in by_first.get(name_core(t[1]).split(" ")[0], []):
        ns = name_sim(t[1], r["Navn"])
        ds = difflib.SequenceMatcher(None, norm(t[2]), norm(r["Beskrivelse"])).ratio()
        sc = (ns + ds) / 2
        if ns >= NAME_MIN_SCORE and ds >= 0.85 and sc > best[0]:
            best, second = (sc, r), best[0]
        elif sc > second:
            second = sc
    return best[1] if best[1] is not None and best[0] - second > 0.05 else None


def load_target():
    ws = load_workbook(TARGET, read_only=True)["Person"]
    return [r for r in ws.iter_rows(min_row=5, values_only=True) if r[0]]


def load_review():
    ws = load_workbook(REVIEW, read_only=True)["Personregister"]
    head = None
    rows = []
    for r in ws.iter_rows(values_only=True):
        if head is None:
            head = list(r)
            continue
        rows.append(dict(zip(head, r)))
    return rows


def match_to_reg(target, entities):
    """V0.97-række -> (Reg-id, kobling, score). Label+beskrivelse, ellers rækkefølge."""
    pos = {e["entity_id"]: i for i, e in enumerate(entities)}
    idx = defaultdict(list)
    for e in entities:
        idx[(norm(e["label"]), norm(e["description"]))].append(e["entity_id"])

    out = {}
    for t in target:
        c = idx.get((norm(t[1]), norm(t[2])), [])
        if len(c) == 1:
            out[t[0]] = (c[0], "label+beskrivelse", 1.0)

    # resten: bedste match i vinduet mellem naboernes Reg-id'er (rækkefølgen er bevaret)
    prev = -1
    for i, t in enumerate(target):
        if t[0] in out:
            prev = pos[out[t[0]][0]]
            continue
        nxt = next((pos[out[x[0]][0]] for x in target[i + 1:] if x[0] in out), len(entities))
        best = None
        for e in entities[prev + 1:nxt]:
            sc = difflib.SequenceMatcher(
                None, norm(t[1]) + "|" + norm(t[2]),
                norm(e["label"]) + "|" + norm(e["description"])).ratio()
            if not best or sc > best[0]:
                best = (sc, e["entity_id"])
        if best:
            out[t[0]] = (best[1], "rækkefølge", best[0])
            prev = pos[best[1]]
    return out


def main():
    target = load_target()
    entities = [r for r in csv.DictReader(open(ENTITIES, encoding="utf-8"))
                if r["entity_type"] == "person"]
    review = load_review()
    by_reg = {r["Reg-ID (live site)"]: r for r in review if r["Reg-ID (live site)"]}
    by_hcap = {r["HCAP-ID"]: r for r in review}
    cross = {r["reg_id"]: r["person_id"]
             for r in csv.DictReader(open(CROSSWALK, encoding="utf-8")) if r["reg_id"]}

    reg_of = match_to_reg(target, entities)
    by_name_desc, by_name, by_first = defaultdict(list), defaultdict(list), defaultdict(list)
    for r in review:
        by_first[name_core(r["Navn"]).split(" ")[0]].append(r)
        by_name_desc[(name_core(r["Navn"]), norm(r["Beskrivelse"]))].append(r)
        by_name[name_core(r["Navn"])].append(r)
    rows, stat = [], Counter()
    for t in target:
        reg, how, score = reg_of.get(t[0], (None, "ingen", 0))
        rv, via = by_reg.get(reg), "Reg-ID"
        if rv is None and reg in cross:
            rv, via = by_hcap.get(cross[reg]), "crosswalk"
        if rv is not None and name_sim(t[1], rv["Navn"]) < NAME_MIN_SCORE:
            rv, via = None, "forkastet (navnene afviger)"   # fejlkobling via Reg-id/rækkefølge
        if rv is None:                                       # reserve: navn + beskrivelse direkte
            c = by_name_desc.get((name_core(t[1]), norm(t[2])), [])
            if len(c) == 1:
                rv, via = c[0], "navn+beskrivelse"
            elif not c and len(by_name.get(name_core(t[1]), [])) == 1:
                rv, via = by_name[name_core(t[1])][0], "kun navn (entydigt)"
            elif not c:                                      # stavevarianter (aa/å, Ernst/Emst, mellemrum)
                rv = fuzzy_match(t, by_first)
                via = "navn+beskrivelse (tilnærmet)"
        flag = []
        if rv is None:
            flag.append("ingen kønsbestemmelse fundet")
        elif via == "kun navn (entydigt)":
            flag.append("koblet alene på navn")
        elif via.endswith("(tilnærmet)"):
            flag.append("tilnærmet kobling på navn+beskrivelse")
        elif name_sim(t[1], rv["Navn"]) < 0.95:
            flag.append("navnene afviger lidt (stavning/OCR) – se efter")
        if how == "rækkefølge" and score < ORDER_MIN_SCORE:
            flag.append(f"svag rækkefølgekobling ({score:.2f})")
        if rv is not None and (rv["Familiegruppe"] or "").strip():
            flag.append("familiegruppe: posten dækker flere personer")
        gender = rv["Køn"] if rv else None
        stat[(gender or "—", "tjek" if flag else "ok")] += 1
        rows.append([
            t[0], t[1], t[15],
            rv["Nationalitet"] if rv else None,
            gender,
            rv["Rolle / Erhverv"] if rv else None,
            rv["Kønsmetode"] if rv else None,
            rv["Kønssikkerhed (0–1)"] if rv else None,
            rv["Kønsgrundlag"] if rv else None,
            rv["Familiegruppe"] if rv else None,
            rv["HCAP-ID"] if rv else None,
            reg,
            f"{how}; {via}" if rv else how,
            "; ".join(flag) or None,
            rv["Navn"] if rv else None,
        ])

    wb = Workbook()
    ws = wb.active
    ws.title = "Gender"
    head = ["PersonID*", "Person*", "OLDPersonID*", "Nationality", "Gender", "Role/Occupation", "Kønsmetode",
            "Kønssikkerhed", "Kønsgrundlag", "Familiegruppe", "HCAP-ID", "Reg-ID",
            "Kobling", "Tjek", "Navn i kønsgennemsyn"]
    ws.append(head)
    for r in rows:
        ws.append(r)
    bold = Font(bold=True)
    fill = PatternFill("solid", fgColor="DDDDDD")
    for c in ws[1]:
        c.font, c.fill = bold, fill
    for i, w in enumerate([13, 50, 13, 14, 38, 34, 30, 12, 60, 14, 11, 12, 28, 40, 50], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions

    info = wb.create_sheet("Læs mig", 0)
    lines = [
        "Kønsbestemmelser til 6-PERSON-V0.97(5)(1).xlsx",
        "",
        "Rækkerne i arket »Gender« står i samme rækkefølge som arket »Person« i V0.97 (række 2 her = række 5 dér).",
        "Kopiér kolonne D:F (Nationality, Gender, Role/Occupation) ind i kolonne H:J i V0.97, fra række 5. Kontrollér først, at PersonID i kolonne A stemmer.",
        "Nationalitet er kun udfyldt, når et nationalitetsadjektiv står først i beskrivelsen; tom = uoplyst (ikke »dansk«). Rolle er parserens rollebuckets (flere adskilt af »; «); tom = ingen rolle fundet.",
        "Kilde: hca-personregister-redesign_gender-review.xlsx, inkl. alle manuelle afgørelser (gender_manual_review.csv). Ingen poster er »Endnu ubestemt«.",
        "Værdier: Mandlig, Kvindelig, Irrelevant (firma, slægt, gruppe, dyr, krydshenvisning: køn er ikke meningsfuldt).",
        "Kønssikkerhed for regel- og modelafgørelser er en sorteringsværdi, ikke en målt præcision. Manuelle afgørelser står som »manuel gennemgang«.",
        "Rækker med »Tjek« udfyldt bør ses efter: svag kobling, manglende kønsbestemmelse eller en post, der samler flere personer (kønnet er da foreløbigt).",
        "Koblingen V0.97 → kønsgennemsynet går via hjemmesidens Reg-id'er (se scripts/enrichment/make_person_gender_transfer.py).",
        "",
        "Optælling (Gender / tjek):",
    ]
    for ln in lines:
        info.append([ln])
    for (g, st), n in sorted(stat.items()):
        info.append([f"  {g}: {n} ({st})"])
    info.append([f"  I alt: {len(rows)}"])
    info["A1"].font = Font(bold=True, size=13)
    info.column_dimensions["A"].width = 150
    for row in info.iter_rows():
        row[0].alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(OUT)
    print("skrev", os.path.relpath(OUT, ROOT))
    for (g, st), n in sorted(stat.items()):
        print(f"  {g:12s} {st:5s} {n}")
    print("  i alt", len(rows))


if __name__ == "__main__":
    main()
