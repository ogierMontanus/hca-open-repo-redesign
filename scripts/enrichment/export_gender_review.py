#!/usr/bin/env python3
"""
export_gender_review.py
-----------------------
Fryser det færdige kønsgennemsyn som en kurateret autoritetstabel på
hjemmesidens id'er (Reg…), så pipelinen kan bruge det uden at læse en
Excel-fil i data/review/:

    data/review/gender_inference/hca-personregister-redesign_gender-review.xlsx
        →  data/curated/person_gender_reviewed.csv

Gennemsynet er lavet over den manuelt rettede segmentering (HCAP-id'er).
Kolonnen »Reg-ID (live site)« kobler hver post til hjemmesidens person. De
619 poster uden Reg-id bruges kun som kandidater til den tilnærmede kobling
nedenfor.

Hjemmesidens personer uden Reg-kobling i gennemsynet får en af to
afgørelser, ellers beholder de parserens værdi:

    krydshenvisning i label (»Aage, se: Drewsen, Aage.«)   Irrelevant, som
        krydshenvisningerne i gennemsynet
    tilnærmet navn + beskrivelse til en gennemset post uden Reg-id
        (stave-/OCR-varianter som »Bayer, Ohr. Fr.« ~ »Bayer, Chr. Fr.«;
        samme regel som make_person_gender_transfer.py)   postens køn

Flere HCAP-poster kan pege på samme Reg-id (hjemmesiden har slået personer
sammen, fx tre forskellige »Bruun« uden fornavn). Er de enige, bruges deres
køn; ellers bliver køn »Endnu ubestemt«, og uenigheden står i grundlaget.

Kør én gang efter et nyt gennemsyn (kræver openpyxl):
    python scripts/enrichment/export_gender_review.py

Tabellen lægges over parserens output af apply_gender_review.py (stage 1f').
"""

import csv
import os
import re
import sys
from collections import Counter, defaultdict

from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from make_person_gender_transfer import fuzzy_match, name_core  # noqa: E402

sys.stdout.reconfigure(encoding="utf-8")
ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
REVIEW = os.path.join(ROOT, "data", "review", "gender_inference",
                      "hca-personregister-redesign_gender-review.xlsx")
ENTITIES = os.path.join(ROOT, "data", "normalized", "entities.csv")
OUT = os.path.join(ROOT, "data", "curated", "person_gender_reviewed.csv")

CROSS_REF = re.compile(r"(^|[\s,])[Ss]e( også)?:?\s")

FIELDS = ["entity_id", "koen", "confidence", "metode", "grundlag",
          "familiegruppe", "hcap_ids"]


def main():
    ws = load_workbook(REVIEW, read_only=True)["Personregister"]
    rows = ws.iter_rows(values_only=True)
    head = next(rows)
    by_reg = defaultdict(list)
    no_reg = defaultdict(list)          # gennemsete poster uden Reg-id, efter første navneord
    for r in rows:
        r = dict(zip(head, r))
        if r["Reg-ID (live site)"]:
            by_reg[r["Reg-ID (live site)"]].append(r)
        else:
            no_reg[name_core(r["Navn"]).split(" ")[0]].append(r)

    extra, how = [], Counter()
    for e in csv.DictReader(open(ENTITIES, encoding="utf-8")):
        if e["entity_type"] != "person" or e["entity_id"] in by_reg:
            continue
        if CROSS_REF.search(e["label"]):
            extra.append([e["entity_id"], "Irrelevant", "", "irrelevant",
                          "krydshenvisning (hjemmesidens label)", "", ""])
            how["krydshenvisning"] += 1
            continue
        r = fuzzy_match((e["entity_id"], e["label"], e["description"]), no_reg)
        if r is not None:
            conf = 1.0 if r["Kønsmetode"] == "manuel gennemgang" else r["Kønssikkerhed (0–1)"]
            extra.append([e["entity_id"], r["Køn"], "" if conf is None else round(float(conf), 3),
                          f"{r['Kønsmetode']} (tilnærmet kobling)",
                          f"koblet til »{r['Navn']}« på navn og beskrivelse; {r['Kønsgrundlag'] or ''}".rstrip("; "),
                          r["Familiegruppe"] or "", r["HCAP-ID"]])
            how["tilnærmet kobling"] += 1

    out = []
    for reg in sorted(by_reg):
        rs = by_reg[reg]
        koens = {r["Køn"] for r in rs}
        ids = " ".join(r["HCAP-ID"] for r in rs)
        fam = " | ".join(r["Familiegruppe"] for r in rs if r["Familiegruppe"])
        if len(koens) == 1:
            r = rs[0]
            conf = r["Kønssikkerhed (0–1)"]
            if r["Kønsmetode"] == "manuel gennemgang":
                conf = 1.0
            metode = r["Kønsmetode"] if len(rs) == 1 else " / ".join(
                dict.fromkeys(x["Kønsmetode"] for x in rs))
            grundlag = r["Kønsgrundlag"] if len(rs) == 1 else (
                f"{len(rs)} poster i registret, alle {r['Køn']}")
            out.append([reg, r["Køn"], "" if conf is None else round(float(conf), 3),
                        metode, grundlag, fam, ids])
        else:
            out.append([reg, "Endnu ubestemt", 0.5, "uenige poster",
                        "hjemmesiden samler flere registerposter med forskelligt køn: "
                        + "; ".join(f"{x['Navn']} = {x['Køn']}" for x in rs),
                        fam, ids])

    out = sorted(out + extra)
    with open(OUT, "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(FIELDS)
        w.writerows(out)
    print(f"skrev {os.path.relpath(OUT, ROOT)}: {len(out):,} personer "
          f"({len(by_reg):,} via Reg-id, " + ", ".join(f"{n} {k}" for k, n in how.items()) + ")")
    for k, n in Counter(r[1] for r in out).most_common():
        print(f"  {k:16s} {n:,}")


if __name__ == "__main__":
    main()
